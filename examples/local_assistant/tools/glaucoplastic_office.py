#!/usr/bin/env python3
"""Cross-platform office tool bridge for GlaucoPlastic RLM.

Invocation:
    python glaucoplastic_office.py <capability> <arguments.json>

The process prints exactly one JSON object to stdout.
"""

from __future__ import annotations

import json
import os
import re
import sys
from pathlib import Path
from typing import Any, Iterable


def _safe_name(value: str, fallback: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "-", value.strip()).strip("-.")
    return cleaned or fallback


def _workspace(arguments: dict[str, Any]) -> Path:
    configured = arguments.get("workspace") or os.environ.get(
        "GLAUCOPLASTIC_OFFICE_OUTPUT", ""
    )
    root = Path(configured).expanduser() if configured else Path.cwd() / "office-output"
    root.mkdir(parents=True, exist_ok=True)
    return root.resolve()


def _inside(root: Path, candidate: Path) -> Path:
    resolved = candidate.expanduser().resolve()
    try:
        resolved.relative_to(root)
    except ValueError as exc:
        raise ValueError(f"Path outside office workspace: {resolved}") from exc
    return resolved


def _output_path(
    arguments: dict[str, Any], extension: str, fallback_stem: str
) -> Path:
    root = _workspace(arguments)
    requested = str(arguments.get("filename") or arguments.get("path") or "").strip()
    if requested:
        candidate = Path(requested)
        if not candidate.is_absolute():
            candidate = root / candidate
    else:
        candidate = root / f"{_safe_name(fallback_stem, 'document')}{extension}"
    if candidate.suffix.lower() != extension:
        candidate = candidate.with_suffix(extension)
    candidate.parent.mkdir(parents=True, exist_ok=True)
    return _inside(root, candidate)


def _paragraph_items(value: Any) -> list[dict[str, str]]:
    if value is None:
        return []
    if isinstance(value, str):
        return [{"text": value}]
    result: list[dict[str, str]] = []
    if isinstance(value, list):
        for item in value:
            if isinstance(item, str):
                result.append({"text": item})
            elif isinstance(item, dict):
                result.append({str(k): str(v) for k, v in item.items() if v is not None})
    return result


def create_document(arguments: dict[str, Any]) -> dict[str, Any]:
    from docx import Document

    title = str(arguments.get("title") or "Documento")
    output = _output_path(arguments, ".docx", title)
    document = Document()
    if title:
        document.add_heading(title, level=0)
    for item in _paragraph_items(arguments.get("paragraphs") or arguments.get("content")):
        heading = item.get("heading") or item.get("title")
        if heading:
            level = int(item.get("level", "1") or 1)
            document.add_heading(heading, level=max(1, min(level, 9)))
        text = item.get("text") or item.get("body") or ""
        if text:
            document.add_paragraph(text)
        bullets = item.get("bullets")
        if isinstance(bullets, list):
            for bullet in bullets:
                document.add_paragraph(str(bullet), style="List Bullet")
    document.save(output)
    return {"ok": True, "kind": "document", "path": str(output), "title": title}


def _normalize_rows(rows: Any) -> list[list[Any]]:
    if not isinstance(rows, list):
        return []
    result: list[list[Any]] = []
    for row in rows:
        if isinstance(row, list):
            result.append(row)
        elif isinstance(row, dict):
            result.append(list(row.values()))
        else:
            result.append([row])
    return result


def create_spreadsheet(arguments: dict[str, Any]) -> dict[str, Any]:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    title = str(arguments.get("title") or "Planilha")
    output = _output_path(arguments, ".xlsx", title)
    workbook = Workbook()
    default = workbook.active
    sheets = arguments.get("sheets")
    if not isinstance(sheets, list) or not sheets:
        sheets = [{"name": arguments.get("sheet") or "Dados", "rows": arguments.get("rows", [])}]
    for index, sheet_data in enumerate(sheets):
        sheet_data = sheet_data if isinstance(sheet_data, dict) else {}
        name = str(sheet_data.get("name") or f"Planilha {index + 1}")[:31]
        worksheet = default if index == 0 else workbook.create_sheet()
        worksheet.title = name
        headers = sheet_data.get("headers")
        if isinstance(headers, list):
            worksheet.append(headers)
            for cell in worksheet[1]:
                cell.font = Font(bold=True)
        for row in _normalize_rows(sheet_data.get("rows")):
            worksheet.append(row)
        for column_cells in worksheet.columns:
            width = min(60, max(10, max(len(str(cell.value or "")) for cell in column_cells) + 2))
            worksheet.column_dimensions[column_cells[0].column_letter].width = width
        worksheet.freeze_panes = "A2" if isinstance(headers, list) and headers else None
    workbook.save(output)
    return {"ok": True, "kind": "spreadsheet", "path": str(output), "title": title}


def create_presentation(arguments: dict[str, Any]) -> dict[str, Any]:
    from pptx import Presentation

    title = str(arguments.get("title") or "Apresentação")
    output = _output_path(arguments, ".pptx", title)
    presentation = Presentation()
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = title
    subtitle = str(arguments.get("subtitle") or "")
    if subtitle and len(title_slide.placeholders) > 1:
        title_slide.placeholders[1].text = subtitle
    slides = arguments.get("slides")
    if not isinstance(slides, list):
        slides = []
    for item in slides:
        item = item if isinstance(item, dict) else {"body": str(item)}
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = str(item.get("title") or "")
        frame = slide.placeholders[1].text_frame
        frame.clear()
        bullets = item.get("bullets")
        if isinstance(bullets, list):
            for index, bullet in enumerate(bullets):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                paragraph.text = str(bullet)
                paragraph.level = 0
        else:
            frame.text = str(item.get("body") or item.get("text") or "")
    presentation.save(output)
    return {"ok": True, "kind": "presentation", "path": str(output), "title": title}


def create_pdf(arguments: dict[str, Any]) -> dict[str, Any]:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    title = str(arguments.get("title") or "Documento")
    output = _output_path(arguments, ".pdf", title)
    styles = getSampleStyleSheet()
    story: list[Any] = [Paragraph(title, styles["Title"]), Spacer(1, 0.5 * cm)]
    for item in _paragraph_items(arguments.get("paragraphs") or arguments.get("content")):
        heading = item.get("heading") or item.get("title")
        if heading:
            story.append(Paragraph(heading, styles["Heading2"]))
        body = item.get("text") or item.get("body") or ""
        if body:
            story.append(Paragraph(body.replace("\n", "<br/>"), styles["BodyText"]))
            story.append(Spacer(1, 0.3 * cm))
    SimpleDocTemplate(str(output), pagesize=A4).build(story)
    return {"ok": True, "kind": "pdf", "path": str(output), "title": title}


def _extract_docx(path: Path) -> str:
    from docx import Document

    document = Document(path)
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def _extract_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    lines: list[str] = []
    for worksheet in workbook.worksheets:
        lines.append(f"# {worksheet.title}")
        for row in worksheet.iter_rows(values_only=True):
            lines.append("\t".join("" if value is None else str(value) for value in row))
    return "\n".join(lines)


def _extract_pptx(path: Path) -> str:
    from pptx import Presentation

    presentation = Presentation(path)
    lines: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        lines.append(f"# Slide {index}")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                lines.append(shape.text)
    return "\n".join(lines)


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)


def extract_text(arguments: dict[str, Any]) -> dict[str, Any]:
    root = _workspace(arguments)
    requested = str(arguments.get("path") or "").strip()
    if not requested:
        raise ValueError("office.text.extract requires path")
    path = Path(requested)
    if not path.is_absolute():
        path = root / path
    path = _inside(root, path)
    if not path.is_file():
        raise FileNotFoundError(path)
    extension = path.suffix.lower()
    if extension == ".docx":
        content = _extract_docx(path)
    elif extension == ".xlsx":
        content = _extract_xlsx(path)
    elif extension == ".pptx":
        content = _extract_pptx(path)
    elif extension == ".pdf":
        content = _extract_pdf(path)
    else:
        content = path.read_text(encoding="utf-8", errors="replace")
    max_chars = int(arguments.get("maxChars") or arguments.get("max_chars") or 30000)
    return {
        "ok": True,
        "kind": "text",
        "path": str(path),
        "content": content[: max(100, max_chars)],
        "truncated": len(content) > max_chars,
    }


def list_files(arguments: dict[str, Any]) -> dict[str, Any]:
    root = _workspace(arguments)
    requested = str(arguments.get("directory") or ".")
    directory = _inside(root, root / requested)
    pattern = str(arguments.get("pattern") or "*")
    recursive = bool(arguments.get("recursive", False))
    iterator: Iterable[Path] = directory.rglob(pattern) if recursive else directory.glob(pattern)
    files = []
    for path in iterator:
        if path.is_file():
            stat = path.stat()
            files.append(
                {
                    "path": str(path),
                    "relativePath": str(path.relative_to(root)),
                    "size": stat.st_size,
                    "extension": path.suffix.lower(),
                }
            )
    files.sort(key=lambda item: item["relativePath"].lower())
    return {"ok": True, "kind": "files", "workspace": str(root), "files": files[:500]}


def workspace_summary(arguments: dict[str, Any]) -> dict[str, Any]:
    listing = list_files({**arguments, "recursive": True, "pattern": "*"})
    totals: dict[str, int] = {}
    total_size = 0
    for item in listing["files"]:
        extension = item["extension"] or "(sem extensão)"
        totals[extension] = totals.get(extension, 0) + 1
        total_size += int(item["size"])
    return {
        "ok": True,
        "kind": "workspace-summary",
        "workspace": listing["workspace"],
        "fileCount": len(listing["files"]),
        "totalSize": total_size,
        "byExtension": totals,
        "files": listing["files"][:100],
    }


CAPABILITIES = {
    "office.document.create": create_document,
    "office.spreadsheet.create": create_spreadsheet,
    "office.presentation.create": create_presentation,
    "office.pdf.create": create_pdf,
    "office.text.extract": extract_text,
    "office.files.list": list_files,
    "office.workspace.summary": workspace_summary,
}


def main() -> int:
    if len(sys.argv) != 3:
        print(json.dumps({"ok": False, "error": "usage: bridge capability arguments.json"}))
        return 2
    capability = sys.argv[1]
    input_path = Path(sys.argv[2])
    try:
        arguments = json.loads(input_path.read_text(encoding="utf-8"))
        if not isinstance(arguments, dict):
            raise TypeError("arguments must be a JSON object")
        handler = CAPABILITIES.get(capability)
        if handler is None:
            raise KeyError(f"Unknown office capability: {capability}")
        result = handler(arguments)
        print(json.dumps(result, ensure_ascii=False))
        return 0
    except Exception as exc:  # The Nim side turns this into a tool failure.
        print(
            json.dumps(
                {"ok": False, "capability": capability, "error": f"{type(exc).__name__}: {exc}"},
                ensure_ascii=False,
            )
        )
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
